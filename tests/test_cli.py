from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from extract import cli
from extract.core.models import Category
from extract.core.pipeline import ScanOptions
from extract.core.progress import Cancelled
from extract.reports import save_report
from test_reports import _machine


def _scan_args(*extra: str):
    return cli.build_parser().parse_args(["scan", *extra])


def test_scan_defaults():
    options = cli.scan_options(_scan_args())
    assert options == ScanOptions()


def test_scan_flags():
    args = _scan_args("--client", " Padaria ", "--quick", "--no-health", "--only", "cpu, storage")
    options = cli.scan_options(args)
    assert options.client == "Padaria"
    assert options.quick_benchmarks
    assert not options.health
    assert options.benchmarks == {Category.CPU, Category.STORAGE}


def test_no_benchmarks_wins_over_only():
    options = cli.scan_options(_scan_args("--no-benchmarks", "--only", "cpu"))
    assert options.benchmarks == set()


@pytest.mark.parametrize("value", ["gpu", "", "cpu,rede"])
def test_only_rejects_unknown_areas(value, capsys):
    with pytest.raises(SystemExit) as exc:
        _scan_args("--only", value)
    assert exc.value.code == 2
    assert "área inválida" in capsys.readouterr().err


def test_command_is_required():
    with pytest.raises(SystemExit):
        cli.build_parser().parse_args([])


def test_scan_saves_json_and_exports(tmp_path, monkeypatch, capsys):
    seen: list[ScanOptions] = []

    def fake_run_scan(options, progress, cancel):
        seen.append(options)
        progress(0.5, "Medindo o disco")
        report = _machine("CAIXA-01", slow_disk=True)
        report.client = options.client
        return report

    monkeypatch.setattr(cli, "run_scan", fake_run_scan)
    output = tmp_path / "padaria"
    code = cli.main(["scan", "--client", "Padaria", "--output", str(output), "--quick"])

    assert code == cli.EXIT_OK
    assert seen[0].quick_benchmarks
    assert len(list(output.glob("CAIXA-01_*.json"))) == 1
    assert (output / "relatorio-padaria.xlsx").exists()
    assert (output / "relatorio-padaria.pdf").exists()
    captured = capsys.readouterr()
    assert "[Crítico] Windows instalado em HD mecânico" in captured.out
    assert "Medindo o disco" in captured.err


def test_scan_json_only_uses_default_folder(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        cli, "run_scan", lambda options, progress, cancel: _machine("PC", slow_disk=False)
    )
    assert cli.main(["scan", "--json-only"]) == cli.EXIT_OK
    folder = tmp_path / "reports" / "sem-cliente"
    assert len(list(folder.glob("*.json"))) == 1
    assert not list(folder.glob("*.xlsx"))


def test_scan_cancelled(tmp_path, monkeypatch, capsys):
    def cancelled(options, progress, cancel):
        raise Cancelled()

    monkeypatch.setattr(cli, "run_scan", cancelled)
    assert cli.main(["scan", "--output", str(tmp_path)]) == cli.EXIT_CANCELLED
    assert "cancelada" in capsys.readouterr().err
    assert not list(tmp_path.iterdir())


def test_scan_error(tmp_path, monkeypatch, capsys):
    def broken(options, progress, cancel):
        raise RuntimeError("falha inesperada")

    monkeypatch.setattr(cli, "run_scan", broken)
    assert cli.main(["scan", "--output", str(tmp_path)]) == cli.EXIT_ERROR
    assert "falha inesperada" in capsys.readouterr().err


def test_export_defaults_to_both_files(tmp_path):
    save_report(_machine("CAIXA-01", slow_disk=True), tmp_path)
    save_report(_machine("ESCRITORIO", slow_disk=False), tmp_path)
    assert cli.main(["export", str(tmp_path)]) == cli.EXIT_OK
    excel = tmp_path / f"relatorio-{tmp_path.name}.xlsx"
    assert load_workbook(excel)["Resumo"].max_row == 3
    assert (tmp_path / f"relatorio-{tmp_path.name}.pdf").read_bytes().startswith(b"%PDF")


def test_export_only_requested_format(tmp_path):
    save_report(_machine("CAIXA-01", slow_disk=True), tmp_path)
    target = tmp_path / "saida" / "cliente.pdf"
    assert cli.main(["export", str(tmp_path), "--pdf", str(target)]) == cli.EXIT_OK
    assert target.exists()
    assert not list(tmp_path.glob("*.xlsx"))


def test_export_errors(tmp_path, capsys):
    assert cli.main(["export", str(tmp_path / "nao-existe")]) == cli.EXIT_ERROR
    assert cli.main(["export", str(tmp_path)]) == cli.EXIT_ERROR
    assert "nenhum relatório" in capsys.readouterr().err


def test_progress_line_skips_repeats_when_redirected():
    stream = io.StringIO()
    line = cli.ProgressLine(stream)
    line(0.1, "Lendo discos")
    line(0.2, "Lendo discos")
    line(0.5, "Medindo CPU")
    assert stream.getvalue() == "[ 10%] Lendo discos\n[ 50%] Medindo CPU\n"
