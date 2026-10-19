from __future__ import annotations

import re
import sys
from datetime import datetime

import pytest
from openpyxl import load_workbook

from extract.core.models import (
    BenchmarkResult,
    Category,
    DiskInfo,
    Finding,
    GpuInfo,
    HealthCheck,
    MachineReport,
    MemoryModule,
    Metric,
    NetworkAdapter,
    Severity,
    Status,
    VolumeInfo,
)
from extract.reports import export_excel, export_pdf
from extract.reports import paths as report_paths
from extract.reports.summary import format_metric, format_number, inventory_rows


def _machine(hostname: str, *, slow_disk: bool) -> MachineReport:
    report = MachineReport(
        client="Padaria Pão & Cia", created_at="2026-10-19T14:30:00-03:00", tool_version="2.0.0"
    )
    inv = report.inventory
    inv.system.hostname = hostname
    inv.system.os_name = "Microsoft Windows 10 Pro"
    inv.system.os_version = "22H2"
    inv.system.os_build = "19045.4291"
    inv.system.boot_mode = "UEFI"
    inv.cpu.name = "Intel(R) Core(TM) i3-7100 CPU @ 3.90GHz"
    inv.cpu.cores, inv.cpu.logical_processors, inv.cpu.base_clock_mhz = 2, 4, 3912
    inv.memory.total_gb, inv.memory.total_slots, inv.memory.memory_type = 4.0, 2, "DDR4"
    inv.memory.modules = [
        MemoryModule("DIMM1", 4.0, 2400, 2133, "Kingston", "KVR24N17S6/4", "DDR4")
    ]
    inv.disks = [
        DiskInfo(
            "WDC WD5000AAKX",
            465.8,
            "HDD" if slow_disk else "SSD",
            "SATA",
            "",
            "Healthy",
            is_system_disk=True,
        )
    ]
    inv.volumes = [VolumeInfo("C:", "", "NTFS", 465.2, 18.4)]
    inv.gpus = [GpuInfo("Intel(R) HD Graphics 630", 1024, "31.0.101.2111", "2022-08-01")]
    inv.network = [NetworkAdapter("Ethernet", "00:11:22:33:44:55", ["192.168.0.10"], dhcp=True)]

    report.health = [
        HealthCheck(
            "storage.free_space",
            Category.STORAGE,
            "Espaço livre",
            Status.CRITICAL,
            "Disco C: com 4% livre.",
        ),
        HealthCheck(
            "system.uptime", Category.SYSTEM, "Tempo ligado", Status.OK, "Ligado há 1 dia."
        ),
        HealthCheck("system.antivirus", Category.SYSTEM, "Antivírus", Status.UNKNOWN, "Sem dados."),
    ]
    report.benchmarks = [
        BenchmarkResult(
            "storage.sequential",
            Category.STORAGE,
            "Disco: leitura e gravação sequencial",
            {
                "read_mb_s": Metric(98.4, "MB/s", "Leitura sequencial"),
                "write_mb_s": Metric(87.12, "MB/s", "Gravação sequencial"),
            },
            duration_s=21.37,
        ),
        BenchmarkResult("cpu.multi", Category.CPU, "Processador", error="Teste cancelado"),
    ]
    if slow_disk:
        report.findings = [
            Finding(
                "storage.hdd_system",
                Category.STORAGE,
                Severity.CRITICAL,
                "Windows instalado em HD mecânico",
                "O disco do sistema é um HD de 7200 rpm <lento> & antigo.",
                "Trocar por um SSD de 240 GB ou mais.",
            ),
            Finding(
                "memory.low",
                Category.MEMORY,
                Severity.WARNING,
                "Pouca memória RAM",
                "4 GB de RAM para o Windows 10.",
                "Adicionar um módulo de 4 GB DDR4.",
            ),
            Finding("info", Category.SYSTEM, Severity.INFO, "Sem dados de antivírus", "", ""),
        ]
    return report


@pytest.fixture
def reports() -> list[MachineReport]:
    return [_machine("CAIXA-01", slow_disk=True), _machine("ESCRITORIO", slow_disk=False)]


def test_format_number_uses_brazilian_separators():
    assert format_number(1234.5) == "1.234,5"
    assert format_number(16.0) == "16"
    assert format_metric(Metric(98.4, "MB/s", "")) == "98,4 MB/s"
    assert format_metric(Metric(0.123, "ms", "")) == "0,12 ms"


def test_inventory_rows_describe_hardware(reports):
    rows = {component: (name, details) for component, name, details in inventory_rows(reports[0])}
    assert rows["Disco do sistema"] == ("WDC WD5000AAKX", "HDD SATA · 465,8 GB · Saudável")
    assert rows["Memória"][0] == "4 GB DDR4 (1 de 2 slots)"
    assert "18,4 GB livres" in rows["Unidade C:"][1]


def test_excel_export(tmp_path, reports):
    path = export_excel(reports, tmp_path / "saida" / "relatorio.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Resumo", "Achados", "Desempenho", "Inventário"]

    summary = workbook["Resumo"]
    header = [c.value for c in summary[1]]
    first = dict(zip(header, (c.value for c in summary[2]), strict=True))
    assert first["Computador"] == "CAIXA-01"
    assert first["Data da análise"] == datetime(2026, 10, 19, 14, 30)
    assert first["RAM total (GB)"] == 4.0
    assert first["Slots (usados/total)"] == "1/2"
    assert first["Disco do sistema"] == "HDD SATA"
    assert first["Principal achado (gravidade)"] == "Crítico"
    assert (first["Críticos"], first["Atenção"]) == (1, 1)
    second = dict(zip(header, (c.value for c in summary[3]), strict=True))
    assert second["Principal achado"] == "Nenhum problema encontrado"
    assert summary.freeze_panes == "B2"
    assert summary.auto_filter.ref == "A1:M3"
    severity_cell = summary.cell(row=2, column=header.index("Principal achado (gravidade)") + 1)
    assert severity_cell.fill.fgColor.rgb.endswith("FDE2E1")

    findings = workbook["Achados"]
    assert findings.max_row == 4
    assert [findings.cell(row=r, column=2).value for r in (2, 3, 4)] == [
        "Crítico",
        "Atenção",
        "Informativo",
    ]

    performance = workbook["Desempenho"]
    rows = list(performance.iter_rows(min_row=2, values_only=True))
    assert rows[0][:2] == ("CAIXA-01", "Disco: leitura e gravação sequencial")
    assert rows[0][4] == 98.4
    assert any(row[-1] == "Teste cancelado" for row in rows)

    inventory = workbook["Inventário"]
    components = [inventory.cell(row=r, column=2).value for r in range(2, inventory.max_row + 1)]
    assert "Processador" in components
    assert "Disco do sistema" in components


def test_excel_export_with_no_reports(tmp_path):
    path = export_excel([], tmp_path / "vazio.xlsx")
    assert load_workbook(path)["Resumo"].max_row == 1


def _page_count(data: bytes) -> int:
    return len(re.findall(rb"/Type\s*/Page(?!s)", data))


def test_pdf_export(tmp_path, reports):
    path = export_pdf(reports, tmp_path / "relatorio.pdf")
    data = path.read_bytes()
    assert data.startswith(b"%PDF")
    assert _page_count(data) >= 3  # cover + one section per machine


def test_pdf_export_single_machine_without_findings(tmp_path):
    report = _machine("SOZINHO", slow_disk=False)
    report.client = ""
    report.health = []
    report.benchmarks = []
    path = export_pdf([report], tmp_path / "um.pdf")
    assert path.read_bytes().startswith(b"%PDF")


def test_default_reports_root_from_source(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.delattr(sys, "frozen", raising=False)
    assert report_paths.client_folder("") == tmp_path / "reports" / "sem-cliente"
    assert report_paths.client_folder("Pão & Cia") == tmp_path / "reports" / "Pao-Cia"


def test_default_reports_root_when_frozen(tmp_path, monkeypatch):
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(sys, "executable", str(tmp_path / "Extract.exe"))
    assert report_paths.app_dir() == tmp_path.resolve()
    assert report_paths.default_reports_root() == tmp_path.resolve() / "Relatorios"


def test_exports_handle_an_empty_report(tmp_path):
    empty = [MachineReport()]
    assert export_pdf(empty, tmp_path / "vazio.pdf").read_bytes().startswith(b"%PDF")
    assert load_workbook(export_excel(empty, tmp_path / "vazio.xlsx"))["Resumo"].max_row == 2


def test_default_reports_root_falls_back_when_read_only(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.delattr(sys, "frozen", raising=False)
    monkeypatch.setattr(report_paths, "_is_writable", lambda folder: False)
    monkeypatch.setattr(report_paths.Path, "home", lambda: tmp_path / "home")
    assert report_paths.default_reports_root() == tmp_path / "home" / "Documents" / "Extract"
