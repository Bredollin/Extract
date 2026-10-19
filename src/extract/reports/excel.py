"""Excel export: one workbook consolidating every machine of a client.

Sheets: "Resumo" (one row per machine), "Achados", "Desempenho" and "Inventário". Each
sheet is a plain table with a styled header, frozen header row and autofilter, so the
technician can sort and filter without reformatting anything.
"""

from __future__ import annotations

import logging
from collections.abc import Sequence
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from extract.core.models import MachineReport, Severity
from extract.reports.summary import (
    SEVERITY_COLORS,
    disk_type,
    inventory_rows,
    os_label,
    severity_counts,
    sorted_findings,
    system_disk,
    top_finding,
)

log = logging.getLogger(__name__)

_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_BORDER = Border(bottom=Side(style="thin", color="D0D7E2"))
_WRAP = Alignment(wrap_text=True, vertical="top")
_TOP = Alignment(vertical="top")
_MAX_WIDTH = 60

# Column headers per sheet; columns listed in _WRAPPED hold sentences and wrap.
_SUMMARY_COLUMNS = [
    "Computador",
    "Cliente",
    "Data da análise",
    "Sistema operacional",
    "Processador",
    "RAM total (GB)",
    "Tipo de RAM",
    "Slots (usados/total)",
    "Disco do sistema",
    "Principal achado (gravidade)",
    "Principal achado",
    "Críticos",
    "Atenção",
]
_FINDING_COLUMNS = [
    "Computador",
    "Gravidade",
    "Área",
    "Achado",
    "Detalhes",
    "Recomendação",
]
_BENCHMARK_COLUMNS = [
    "Computador",
    "Teste",
    "Área",
    "Métrica",
    "Valor",
    "Unidade",
    "Duração (s)",
    "Erro",
]
_INVENTORY_COLUMNS = ["Computador", "Componente", "Modelo / nome", "Detalhes"]
_WRAPPED = {"Detalhes", "Recomendação", "Principal achado", "Erro"}


def export_excel(reports: Sequence[MachineReport], path: Path) -> Path:
    """Write the consolidated workbook to `path` and return it."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo"
    _fill(summary, _SUMMARY_COLUMNS, [_summary_row(r) for r in reports])
    _color_severity(summary, _SUMMARY_COLUMNS.index("Principal achado (gravidade)"))

    findings = workbook.create_sheet("Achados")
    _fill(findings, _FINDING_COLUMNS, [row for r in reports for row in _finding_rows(r)])
    _color_severity(findings, _FINDING_COLUMNS.index("Gravidade"))

    benchmarks = workbook.create_sheet("Desempenho")
    _fill(benchmarks, _BENCHMARK_COLUMNS, [row for r in reports for row in _benchmark_rows(r)])

    inventory = workbook.create_sheet("Inventário")
    rows = [(r.hostname, *row) for r in reports for row in inventory_rows(r)]
    _fill(inventory, _INVENTORY_COLUMNS, rows)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    log.info("Excel report saved to %s", path)
    return path


def _summary_row(report: MachineReport) -> list[Any]:
    memory = report.inventory.memory
    counts = severity_counts(report)
    top = top_finding(report)
    return [
        report.hostname,
        report.client,
        _parse_date(report.created_at),
        os_label(report.inventory.system),
        report.inventory.cpu.name,
        memory.total_gb or None,
        memory.memory_type,
        f"{memory.used_slots}/{memory.total_slots}" if memory.total_slots else "",
        disk_type(system_disk(report)),
        top.severity.label if top else "",
        top.title if top else "Nenhum problema encontrado",
        counts[Severity.CRITICAL],
        counts[Severity.WARNING],
    ]


def _finding_rows(report: MachineReport) -> list[list[Any]]:
    return [
        [
            report.hostname,
            f.severity.label,
            f.category.label,
            f.title,
            f.detail,
            f.recommendation,
        ]
        for f in sorted_findings(report)
    ]


def _benchmark_rows(report: MachineReport) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for bench in report.benchmarks:
        duration = round(bench.duration_s, 1) if bench.duration_s else None
        if not bench.metrics:
            label = bench.category.label
            rows.append([report.hostname, bench.title, label, "", None, "", duration, bench.error])
        for metric in bench.metrics.values():
            rows.append(
                [
                    report.hostname,
                    bench.title,
                    bench.category.label,
                    metric.label,
                    metric.value,
                    metric.unit,
                    duration,
                    bench.error,
                ]
            )
    return rows


def _parse_date(iso: str) -> datetime | str:
    """Excel has no time zones: keep the local wall-clock time of the scan."""
    try:
        return datetime.fromisoformat(iso).replace(tzinfo=None)
    except ValueError:
        return iso


def _fill(sheet: Worksheet, headers: list[str], rows: Sequence[Sequence[Any]]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        sheet.append(list(row))

    for index, header in enumerate(headers, start=1):
        alignment = _WRAP if header in _WRAPPED else _TOP
        for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
            cell.alignment = alignment
            cell.border = _BORDER
            if cell.data_type == "f":
                # Report text can come from the client's machine (device names, startup
                # entries) or from an edited JSON file: never let it become a live formula.
                cell.data_type = "s"
                cell.quotePrefix = True
            elif isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
            elif isinstance(cell.value, float):
                cell.number_format = "#,##0.0"
        sheet.column_dimensions[get_column_letter(index)].width = _width(sheet, index, header)

    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions


def _width(sheet: Worksheet, column: int, header: str) -> float:
    longest = len(header)
    for (value,) in sheet.iter_rows(min_row=2, min_col=column, max_col=column, values_only=True):
        if isinstance(value, datetime):
            longest = max(longest, 16)
        elif value is not None:
            longest = max(longest, len(str(value)))
    return min(longest + 3, _MAX_WIDTH)


def _color_severity(sheet: Worksheet, column_index: int) -> None:
    colors = {severity.label: SEVERITY_COLORS[severity] for severity in Severity}
    for (cell,) in sheet.iter_rows(min_row=2, min_col=column_index + 1, max_col=column_index + 1):
        if cell.value in colors:
            text, background = colors[cell.value]
            cell.fill = PatternFill("solid", fgColor=background)
            cell.font = Font(bold=True, color=text)
