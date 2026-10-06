# ==============================================================================
# Script: gui_extract.py
# Interface Gráfica e Coletor Nativo Integrados (Extract T.I. - Standalone)
# ==============================================================================

import os
import sys
import glob
import re
import platform
import subprocess
import json
import shutil
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog

try:
    import ttkbootstrap as ttk
    from ttkbootstrap.constants import *
except ImportError:
    from tkinter import ttk
    PRIMARY = "primary"
    SECONDARY = "secondary"
    SUCCESS = "success"
    INFO = "info"
    WARNING = "warning"

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Imports do ReportLab para geração do PDF
try:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False


# --- MÓDULO DE COLETA NATIVA DE HARDWARE (EXTRAC INTEGRADO) ---

def run_ps_native(cmd):
    """Executa comando PowerShell em JSON e retorna dicionário/lista."""
    ps_command = f"powershell -NoProfile -Command \"{cmd} | ConvertTo-Json -Depth 3\""
    try:
        res = subprocess.run(ps_command, capture_output=True, text=True, shell=True, encoding='utf-8', errors='ignore')
        if not res.stdout.strip():
            return []
        data = json.loads(res.stdout)
        return [data] if isinstance(data, dict) else data
    except Exception:
        return []

def collect_system_native(log_func):
    log_func("[EXTRAC] Coletando informações gerais do Sistema Operacional e Máquina...")
    os_info = run_ps_native("Get-CimInstance Win32_OperatingSystem")
    cs_info = run_ps_native("Get-CimInstance Win32_ComputerSystem")
    bb_info = run_ps_native("Get-CimInstance Win32_BIOS")
    
    os_data = os_info[0] if os_info else {}
    cs_data = cs_info[0] if cs_info else {}
    bb_data = bb_info[0] if bb_info else {}

    return {
        "Nome da Máquina": platform.node(),
        "Usuário Logado": cs_data.get("UserName", "N/A"),
        "Sistema Operacional": os_data.get("Caption", platform.system()),
        "Arquitetura": os_data.get("OSArchitecture", "N/A"),
        "Versão / Build": os_data.get("Version", "N/A"),
        "Fabricante da Maquina": cs_data.get("Manufacturer", "N/A"),
        "Modelo da Maquina": cs_data.get("Model", "N/A"),
        "Modo Boot / BIOS": bb_data.get("SMBIOSBIOSVersion", "N/A")
    }

def collect_cpu_native(log_func):
    log_func("[EXTRAC] Identificando Processador (CPU), núcleos e soquete...")
    cpu_info = run_ps_native("Get-CimInstance Win32_Processor")
    if not cpu_info:
        return {"Processador": platform.processor()}
    
    c = cpu_info[0]
    return {
        "Modelo CPU": c.get("Name", "").strip(),
        "Núcleos Físicos": c.get("NumberOfCores", "N/A"),
        "Processadores Lógicos": c.get("NumberOfLogicalProcessors", "N/A"),
        "Frequência Base": f"{c.get('MaxClockSpeed', 0)} MHz",
        "Soquete / Soquete": c.get("SocketDesignation", "N/A")
    }

def collect_motherboard_native(log_func):
    log_func("[EXTRAC] Coletando dados da Placa-Mãe e BIOS...")
    mb = run_ps_native("Get-CimInstance Win32_BaseBoard")
    mb_data = mb[0] if mb else {}
    return {
        "Fabricante Placa-Mãe": mb_data.get("Manufacturer", "N/A"),
        "Modelo Placa-Mãe": mb_data.get("Product", "N/A"),
        "Número de Série": mb_data.get("SerialNumber", "N/A")
    }

def collect_ram_native(log_func):
    log_func("[EXTRAC] Analisando Memória RAM (GB total, slots ocupados, frequência e geração DDR)...")
    ram_info = run_ps_native("Get-CimInstance Win32_PhysicalMemory")
    ram_array = run_ps_native("Get-CimInstance Win32_PhysicalMemoryArray")
    
    total_slots = ram_array[0].get("MemoryDevices", 2) if ram_array else 2
    total_bytes = sum([int(r.get("Capacity", 0)) for r in ram_info])
    total_gb = total_bytes // (1024**3)

    pentes = []
    freqs = []
    smbios_types = []

    for idx, r in enumerate(ram_info, 1):
        cap = int(r.get("Capacity", 0)) // (1024**3)
        spd = r.get("Speed", 0)
        mfg = r.get("Manufacturer", "Genérica").strip()
        part = r.get("PartNumber", "").strip()
        loc = r.get("DeviceLocator", f"Slot {idx}").strip()
        smbios_types.append(r.get("SMBIOSMemoryType", 0))
        freqs.append(spd)
        pentes.append(f"[{loc}] -> {cap} GB {mfg} ({spd} MHz) | Part: {part}")

    ddr = "DDR4"
    if 34 in smbios_types or any(f >= 4800 for f in freqs): ddr = "DDR5"
    elif 26 in smbios_types or any(f >= 2133 for f in freqs): ddr = "DDR4"
    elif 24 in smbios_types or any(f < 2133 for f in freqs if f > 0): ddr = "DDR3"

    return {
        "Total RAM": f"{total_gb} GB",
        "Geração": ddr,
        "Slots Totais / Ocupados": f"{total_slots} Slots ({len(ram_info)} Ocupados)",
        "Pentes Instalados": pentes
    }

def collect_disks_native(log_func):
    log_func("[EXTRAC] Varrendo unidades de Armazenamento (SSD NVMe, SATA e HDD)...")
    disks = run_ps_native("Get-CimInstance Win32_DiskDrive")
    disk_list = []
    for d in disks:
        mod = d.get("Model", "Desconhecido").strip()
        sz = int(d.get("Size", 0)) // (1024**3)
        media = "SSD NVMe" if "NVMe" in mod else ("SSD" if "SSD" in mod else "HDD / SSD")
        disk_list.append(f"• {mod} ({sz} GB) [{media}] - Interface: {d.get('InterfaceType', 'N/A')}")
    return {"Discos": disk_list if disk_list else ["N/A"]}

def collect_gpu_native(log_func):
    log_func("[EXTRAC] Verificando Placa de Vídeo e VRAM...")
    gpus = run_ps_native("Get-CimInstance Win32_VideoController")
    gpu_list = []
    for g in gpus:
        name = g.get("Name", "").strip()
        if name and "Virtual" not in name:
            ram_mb = int(g.get("AdapterRAM", 0)) // (1024**2) if g.get("AdapterRAM") else 0
            gpu_list.append(f"{name} ({ram_mb} MB VRAM)")
    return {"Placa de Vídeo": ", ".join(gpu_list) if gpu_list else "Vídeo Integrado"}

def collect_network_native(log_func):
    log_func("[EXTRAC] Mapeando Adaptadores de Rede, IPs ativos e Endereços MAC...")
    nets = run_ps_native("Get-CimInstance Win32_NetworkAdapterConfiguration")
    net_list = []
    for n in nets:
        if n.get("IPEnabled", False):
            desc = n.get("Description", "Adaptador")
            ips = n.get("IPAddress", [])
            ip = ips[0] if ips else "N/A"
            mac = n.get("MACAddress", "N/A")
            net_list.append(f"{desc}\n -> IP: {ip} | MAC: {mac}")
    return {"Redes": net_list if net_list else ["N/A"]}

def generate_report_native(options, output_dir, log_func):
    log_func("[EXTRAC] Iniciando varredura modular de hardware...")
    
    report_lines = [
        "==================================================",
        "          RELATÓRIO DE HARDWARE - EXTRACT         ",
        "=================================================="
    ]

    sys_data = collect_system_native(log_func)
    report_lines.append(f"Nome da Maquina: {sys_data['Nome da Máquina']}")
    report_lines.append(f"Sistema Operacional: {sys_data['Sistema Operacional']} {sys_data['Arquitetura']}")
    report_lines.append(f"Usuário: {sys_data['Usuário Logado']}")
    report_lines.append(f"Modelo: {sys_data['Fabricante da Maquina']} - {sys_data['Modelo da Maquina']}")
    report_lines.append("")

    if options.get("all") or options.get("cpu"):
        report_lines.append("=== PROCESSADOR (CPU) ===")
        cpu = collect_cpu_native(log_func)
        for k, v in cpu.items(): report_lines.append(f"{k}: {v}")
        report_lines.append("")

    if options.get("all") or options.get("ram"):
        report_lines.append("=== MEMORIA RAM E SLOTS ===")
        ram = collect_ram_native(log_func)
        report_lines.append(f"Total de RAM Instalada: {ram['Total RAM']}")
        report_lines.append(f"Slots de Memoria: {ram['Slots Totais / Ocupados']}")
        report_lines.append(f"Geracao Detectada: {ram['Geração']}")
        report_lines.append("Pentes Instalados:")
        for p in ram["Pentes Instalados"]: report_lines.append(f"  {p}")
        report_lines.append("")

    if options.get("all") or options.get("mb"):
        report_lines.append("=== PLACA-MAE ===")
        mb = collect_motherboard_native(log_func)
        report_lines.append(f"Fabricante e Modelo: {mb['Fabricante Placa-Mãe']} {mb['Modelo Placa-Mãe']}")
        report_lines.append(f"Número de Série: {mb['Número de Série']}")
        report_lines.append("")

    if options.get("all") or options.get("gpu"):
        report_lines.append("=== PLACA DE VIDEO (GPU) ===")
        gpu = collect_gpu_native(log_func)
        report_lines.append(f"Nome / Modelo: {gpu['Placa de Vídeo']}")
        report_lines.append("")

    if options.get("all") or options.get("disk"):
        report_lines.append("=== ARMAZENAMENTO (DISCOS) ===")
        disks = collect_disks_native(log_func)
        for d in disks["Discos"]: report_lines.append(d)
        report_lines.append("")

    if options.get("all") or options.get("net"):
        report_lines.append("INFORMACOES DE REDE")
        report_lines.append("---------------------------------------------------")
        nets = collect_network_native(log_func)
        for n in nets["Redes"]: report_lines.append(f"Adaptador: {n}")

    report_lines.append("==================================================")

    log_func("[EXTRAC] Salvando relatório no arquivo local...")
    filename = f"{sys_data['Nome da Máquina']}-info.txt"
    filepath = os.path.join(output_dir, filename)
    with open(filepath, "w", encoding="utf-8") as f:
        f.write("\n".join(report_lines))
    
    log_func(f"[EXTRAC] Relatório gerado com sucesso: {filename}")
    return filepath


# --- PARSER DOS ARQUIVOS DE TEXTO ---

def parse_txt_file(file_path):
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        content = f.read()

    def get_val(pattern, text):
        match = re.search(pattern, text, re.IGNORECASE)
        return match.group(1).strip() if match else 'N/A'

    data = {
        'Computador': get_val(r'Nome da Maquina:\s*(.*)', content),
        'Sistema Operacional': get_val(r'Sistema Operacional:\s*(.*)', content),
        'Processador': 'N/A',
        'RAM Total': 'N/A',
        'Tipo / Slots e Pentes RAM': 'N/A',
        'Armazenamento': 'N/A',
        'Placa de Vídeo': 'N/A',
        'Rede / IP / MAC': 'N/A',
        'Placa-Mãe': 'N/A'
    }

    cpu_match = re.search(r'=== PROCESSADOR \(CPU\) ===(.*?)===', content, re.DOTALL)
    if cpu_match:
        data['Processador'] = get_val(r'Modelo CPU:\s*(.*)', cpu_match.group(1))
        if data['Processador'] == 'N/A':
            data['Processador'] = get_val(r'Nome / Modelo:\s*(.*)', cpu_match.group(1))

    ram_match = re.search(r'=== MEMORIA RAM E SLOTS ===(.*?)===', content, re.DOTALL)
    if ram_match:
        ram_text = ram_match.group(1)
        data['RAM Total'] = get_val(r'Total de RAM Instalada:\s*(.*)', ram_text)
        slots = get_val(r'Slots de Memoria:\s*(.*)', ram_text)
        geracao = get_val(r'Geracao Detectada:\s*(.*)', ram_text)
        pentes = re.findall(r'(\[.*?\]\s*->\s*.*)', ram_text)
        pentes_str = " | ".join([p.strip() for p in pentes]) if pentes else "N/A"
        data['Tipo / Slots e Pentes RAM'] = f"[{geracao}] Slots: {slots}\n-> {pentes_str}"

    mb_match = re.search(r'=== PLACA-MAE ===(.*?)===', content, re.DOTALL)
    if mb_match:
        data['Placa-Mãe'] = get_val(r'Fabricante e Modelo:\s*(.*)', mb_match.group(1))

    gpu_match = re.search(r'=== PLACA DE VIDEO \(GPU\) ===(.*?)===', content, re.DOTALL)
    if gpu_match:
        data['Placa de Vídeo'] = get_val(r'Nome / Modelo:\s*(.*)', gpu_match.group(1))

    disk_match = re.search(r'=== ARMAZENAMENTO \(DISCOS\) ===(.*?)INFORMACOES DE REDE', content, re.DOTALL)
    if disk_match:
        disks = re.findall(r'(• .*)', disk_match.group(1))
        data['Armazenamento'] = "\n".join(disks) if disks else 'N/A'

    net_match = re.search(r'INFORMACOES DE REDE\s*---------------------------------------------------\s*(.*)', content, re.DOTALL)
    if net_match:
        adapters = re.findall(r'Adaptador:\s*(.*?)(?=(?:Adaptador:|$))', net_match.group(1), re.DOTALL)
        clean_adapters = [a.strip() for a in adapters if a.strip()]
        data['Rede / IP / MAC'] = "\n".join(clean_adapters) if clean_adapters else 'N/A'

    return data


# --- PROCESSAMENTO EXCEL ---

def process_inventory(input_folder, output_filepath, log_func, progress_callback=None):
    search_pattern = os.path.join(input_folder, "*-info.txt")
    txt_files = glob.glob(search_pattern)

    if not txt_files:
        log_func(f"[ERRO] Nenhum arquivo *-info.txt encontrado em: {input_folder}")
        return False

    log_func(f"[INFO] Processando {len(txt_files)} relatório(s) para Excel...")
    
    data_list = []
    total_files = len(txt_files)
    
    for idx, f in enumerate(txt_files, 1):
        data_list.append(parse_txt_file(f))
        if progress_callback:
            progress_callback((idx / total_files) * 50)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Parque_de_TI"

    header_fill = PatternFill(start_color="1B2A47", end_color="1B2A47", fill_type="solid")
    zebra_fill = PatternFill(start_color="F4F6F9", end_color="F4F6F9", fill_type="solid")
    white_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    regular_font = Font(name="Segoe UI", size=10)
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    border_side = Side(border_style="thin", color="E0E0E0")
    border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    keys = ['Computador', 'Sistema Operacional', 'Processador', 'RAM Total', 
            'Tipo / Slots e Pentes RAM', 'Armazenamento', 'Placa de Vídeo', 'Rede / IP / MAC', 'Placa-Mãe']

    for col_idx, k in enumerate(keys, 1):
        cell = ws.cell(row=1, column=col_idx, value=k)
        cell.font = white_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[1].height = 28

    for r_idx, item in enumerate(data_list, 2):
        is_even = (r_idx % 2 == 0)
        for c_idx, k in enumerate(keys, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=str(item.get(k, 'N/A')))
            cell.font = bold_font if c_idx == 1 else regular_font
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = border
            if is_even: cell.fill = zebra_fill

    widths = [20, 25, 30, 14, 42, 35, 25, 40, 25]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.auto_filter.ref = f"A1:{get_column_letter(len(keys))}{len(data_list) + 1}"
    wb.save(output_filepath)
    log_func(f"[SUCESSO] Planilha Excel salva em:\n ➔ {output_filepath}")
    return True


# --- PROCESSAMENTO E GERAÇÃO DO PDF ---

def process_pdf_report(input_folder, output_pdf_filepath, log_func, progress_callback=None):
    if not HAS_REPORTLAB:
        log_func("[ERRO PDF] Biblioteca 'reportlab' não instalada!")
        return False

    try:
        search_pattern = os.path.join(input_folder, "*-info.txt")
        txt_files = glob.glob(search_pattern)

        if not txt_files:
            log_func(f"[ERRO PDF] Nenhum arquivo *-info.txt encontrado para gerar PDF.")
            return False

        log_func(f"[INFO] Gerando Relatório Executivo PDF...")
        data_list = [parse_txt_file(f) for f in txt_files]

        # Configuração de documento A4 Landscape com margens de 20pt
        doc = SimpleDocTemplate(
            output_pdf_filepath,
            pagesize=landscape(A4),
            rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20
        )

        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=16,
            textColor=colors.HexColor('#1B2A47'),
            spaceAfter=4
        )

        subtitle_style = ParagraphStyle(
            'SubtitleStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            textColor=colors.HexColor('#555555'),
            spaceAfter=12
        )

        cell_style = ParagraphStyle(
            'CellStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8,
            leading=10
        )

        cell_header_style = ParagraphStyle(
            'CellHeaderStyle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=9,
            leading=11,
            textColor=colors.white
        )

        elements = []

        nome_cliente = os.path.basename(input_folder.rstrip(os.sep))
        elements.append(Paragraph(f"Relatório de Inventário de T.I. - {nome_cliente}", title_style))
        elements.append(Paragraph(f"Total de Dispositivos Auditados: {len(data_list)} | Gerado por Extract T.I.", subtitle_style))

        headers = ['Computador', 'Sistema Operacional', 'Processador', 'RAM', 'Armazenamento', 'Rede / IP']
        table_data = [[Paragraph(h, cell_header_style) for h in headers]]

        for idx, item in enumerate(data_list, 1):
            arm_text = item['Armazenamento'].replace('\n', '<br/>')
            net_text = item['Rede / IP / MAC'].replace('\n', '<br/>')

            row = [
                Paragraph(f"<b>{item['Computador']}</b>", cell_style),
                Paragraph(item['Sistema Operacional'], cell_style),
                Paragraph(item['Processador'], cell_style),
                Paragraph(f"{item['RAM Total']}<br/>{item['Tipo / Slots e Pentes RAM']}", cell_style),
                Paragraph(arm_text, cell_style),
                Paragraph(net_text, cell_style),
            ]
            table_data.append(row)
            if progress_callback:
                progress_callback(50 + (idx / len(data_list)) * 50)

        # Soma das colunas = 90 + 110 + 130 + 110 + 140 + 190 = 770pt (Perfeito no A4 Landscape)
        col_widths = [90, 110, 130, 110, 140, 190]

        pdf_table = Table(table_data, colWidths=col_widths, repeatRows=1)
        pdf_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1B2A47')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E0E0E0')),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')])
        ]))

        elements.append(pdf_table)
        doc.build(elements)

        log_func(f"[SUCESSO] Relatório PDF gerado em:\n ➔ {output_pdf_filepath}")
        return True

    except Exception as e:
        log_func(f"[ERRO CRÍTICO PDF] Falha ao criar PDF: {str(e)}")
        return False


# --- INTERFACE GRÁFICA MODERNA ---

class ExtractGUIApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Extract - Diagnóstico & Inventário de T.I.")
        self.root.geometry("820x860")
        self.root.resizable(False, False)

        header_frame = ttk.Frame(root, padding=12)
        header_frame.pack(fill="x")

        title_lbl = ttk.Label(
            header_frame, 
            text="💻 EXTRACT", 
            font=("Segoe UI", 20, "bold"), 
            bootstyle="primary"
        )
        title_lbl.pack(anchor="w")

        sub_lbl = ttk.Label(
            header_frame, 
            text="Central de Diagnóstico de Hardware e Consolidação de Inventário", 
            font=("Segoe UI", 10), 
            bootstyle="secondary"
        )
        sub_lbl.pack(anchor="w")

        # 1. CARD: SELEÇÃO DE CLIENTE / PASTA
        card_cliente = ttk.Labelframe(root, text=" 1. Cliente & Destino dos Dados ", padding=10)
        card_cliente.pack(fill="x", padx=20, pady=4)

        self.entry_cliente_folder = ttk.Entry(card_cliente, font=("Segoe UI", 10))
        self.entry_cliente_folder.insert(0, os.getcwd())
        self.entry_cliente_folder.pack(side="left", fill="x", expand=True, padx=(0, 8))

        btn_sel = ttk.Button(card_cliente, text="📂 Selecionar", bootstyle="outline-primary", command=self.select_client_folder)
        btn_sel.pack(side="left", padx=3)

        btn_novo = ttk.Button(card_cliente, text="➕ Novo Cliente", bootstyle="success", command=self.create_new_client_folder)
        btn_novo.pack(side="left", padx=3)

        # 2. CARD: MÓDULOS DE VARREDURA
        card_opts = ttk.Labelframe(root, text=" 2. Módulos de Coleta da Máquina Atual ", padding=10)
        card_opts.pack(fill="x", padx=20, pady=4)

        self.var_all = tk.BooleanVar(value=True)
        self.var_cpu = tk.BooleanVar(value=True)
        self.var_ram = tk.BooleanVar(value=True)
        self.var_mb  = tk.BooleanVar(value=True)
        self.var_gpu = tk.BooleanVar(value=True)
        self.var_disk= tk.BooleanVar(value=True)
        self.var_net = tk.BooleanVar(value=True)

        chk_all = ttk.Checkbutton(
            card_opts, text="🌟 RELATÓRIO GERAL COMPLETO", 
            variable=self.var_all, bootstyle="round-toggle-primary", command=self.toggle_all
        )
        chk_all.grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 6))

        ttk.Checkbutton(card_opts, text="Processador (CPU)", variable=self.var_cpu, bootstyle="square-primary").grid(row=1, column=0, sticky="w", padx=5, pady=2)
        ttk.Checkbutton(card_opts, text="Memória RAM & Slots", variable=self.var_ram, bootstyle="square-primary").grid(row=1, column=1, sticky="w", padx=5, pady=2)
        ttk.Checkbutton(card_opts, text="Placa-Mãe & BIOS", variable=self.var_mb, bootstyle="square-primary").grid(row=1, column=2, sticky="w", padx=5, pady=2)
        ttk.Checkbutton(card_opts, text="Placa de Vídeo (GPU)", variable=self.var_gpu, bootstyle="square-primary").grid(row=2, column=0, sticky="w", padx=5, pady=2)
        ttk.Checkbutton(card_opts, text="Discos (HD/SSD)", variable=self.var_disk, bootstyle="square-primary").grid(row=2, column=1, sticky="w", padx=5, pady=2)
        ttk.Checkbutton(card_opts, text="Rede & IP/MAC", variable=self.var_net, bootstyle="square-primary").grid(row=2, column=2, sticky="w", padx=5, pady=2)

        self.btn_coletar = ttk.Button(
            card_opts, text="⚡ EXECUTAR VARREDURA NESTA MÁQUINA", 
            bootstyle="warning", command=self.start_coleta_thread
        )
        self.btn_coletar.grid(row=3, column=0, columnspan=3, sticky="ew", pady=(8, 0))

        # BARRA DE PROGRESSO
        self.progress_frame = ttk.Frame(root, padding=(20, 2))
        self.progress_frame.pack(fill="x")

        self.lbl_status_progress = ttk.Label(self.progress_frame, text="Aguardando ação...", font=("Segoe UI", 9), bootstyle="secondary")
        self.lbl_status_progress.pack(anchor="w", pady=(0, 2))

        self.progress_bar = ttk.Progressbar(self.progress_frame, bootstyle="primary-striped", mode="determinate")
        self.progress_bar.pack(fill="x")

        # 3. CARD: CONSOLIDAÇÃO EXCEL & PDF
        card_excel = ttk.Labelframe(root, text=" 3. Consolidação e Exportação de Relatórios ", padding=10)
        card_excel.pack(fill="x", padx=20, pady=4)

        self.var_gen_excel = tk.BooleanVar(value=True)
        self.var_gen_pdf = tk.BooleanVar(value=True)

        f_export_types = ttk.Frame(card_excel)
        f_export_types.pack(fill="x", pady=(0, 6))

        ttk.Checkbutton(f_export_types, text="📊 Planilha Excel (.xlsx)", variable=self.var_gen_excel, bootstyle="square-success").pack(side="left", padx=(0, 15))
        ttk.Checkbutton(f_export_types, text="📄 Relatório Executivo em PDF (.pdf)", variable=self.var_gen_pdf, bootstyle="square-danger").pack(side="left")

        f_path_row = ttk.Frame(card_excel)
        f_path_row.pack(fill="x", pady=(0, 8))

        self.entry_excel_path = ttk.Entry(f_path_row, font=("Segoe UI", 10))
        self.update_excel_default_path(os.getcwd())
        self.entry_excel_path.pack(side="left", fill="x", expand=True, padx=(0, 8))

        btn_excel_dest = ttk.Button(f_path_row, text="💾 Salvar como...", bootstyle="outline-secondary", command=self.select_excel_destination)
        btn_excel_dest.pack(side="left")

        self.btn_run = ttk.Button(
            card_excel, text="🚀 GERAR / ATUALIZAR CONSOLIDAÇÃO DE RELATÓRIOS", 
            bootstyle="primary", command=self.start_process_thread
        )
        self.btn_run.pack(fill="x", pady=(2, 0))

        # LOGS DE SISTEMA
        card_log = ttk.Labelframe(root, text=" Terminal de Logs de Extração e Execução ", padding=8)
        card_log.pack(fill="both", expand=True, padx=20, pady=(4, 15))

        self.txt_log = tk.Text(card_log, height=6, state='disabled', font=("Consolas", 9), bg="#1E1E1E", fg="#00FF66")
        self.txt_log.pack(fill="both", expand=True)

        self.log("[SISTEMA] Módulo pronto e aguardando ações.")
        self.check_files_in_current_folder()

    def toggle_all(self):
        val = self.var_all.get()
        self.var_cpu.set(val)
        self.var_ram.set(val)
        self.var_mb.set(val)
        self.var_gpu.set(val)
        self.var_disk.set(val)
        self.var_net.set(val)

    def log(self, msg):
        self.txt_log.config(state='normal')
        self.txt_log.insert(tk.END, msg + "\n")
        self.txt_log.see(tk.END)
        self.txt_log.config(state='disabled')

    def set_progress(self, val, text=""):
        self.progress_bar['value'] = val
        if text:
            self.lbl_status_progress.config(text=text)
        self.root.update_idletasks()

    def update_excel_default_path(self, folder_path):
        nome_pasta = os.path.basename(folder_path.rstrip(os.sep))
        nome_excel = f"Parque_TI_{nome_pasta}.xlsx" if nome_pasta else "Parque_TI_Cliente.xlsx"
        self.entry_excel_path.delete(0, tk.END)
        self.entry_excel_path.insert(0, os.path.join(folder_path, nome_excel))

    def select_client_folder(self):
        folder = filedialog.askdirectory(title="Selecione a Pasta do Cliente")
        if folder:
            self.entry_cliente_folder.delete(0, tk.END)
            self.entry_cliente_folder.insert(0, folder)
            self.update_excel_default_path(folder)
            self.check_files_in_current_folder()

    def create_new_client_folder(self):
        nome_cliente = simpledialog.askstring("Novo Cliente", "Digite o nome da pasta do cliente:")
        if nome_cliente:
            limpo = re.sub(r'[\\/*?:"<>|]', "", nome_cliente).strip()
            caminho = os.path.join(os.getcwd(), limpo)
            os.makedirs(caminho, exist_ok=True)
            self.entry_cliente_folder.delete(0, tk.END)
            self.entry_cliente_folder.insert(0, caminho)
            self.update_excel_default_path(caminho)
            self.check_files_in_current_folder()

    # --- EXECUÇÃO DE COLETA NATIVA EM THREAD (AUTÔNOMA) ---

    def start_coleta_thread(self):
        threading.Thread(target=self.run_coleta, daemon=True).start()

    def run_coleta(self):
        cliente_folder = self.entry_cliente_folder.get().strip()
        if not os.path.exists(cliente_folder):
            messagebox.showerror("Erro", "A pasta do cliente selecionada não existe.")
            return

        self.btn_coletar.config(state='disabled')
        self.btn_run.config(state='disabled')
        self.set_progress(10, "⚡ Executando varredura nativa de hardware...")
        self.progress_bar.config(mode="indeterminate")
        self.progress_bar.start(10)

        self.log("--------------------------------------------------")
        self.log("⚡ Executando varredura e gerando logs em tempo real...")

        opts = {
            "all": self.var_all.get(),
            "cpu": self.var_cpu.get(),
            "ram": self.var_ram.get(),
            "mb": self.var_mb.get(),
            "gpu": self.var_gpu.get(),
            "disk": self.var_disk.get(),
            "net": self.var_net.get()
        }

        try:
            filepath = generate_report_native(opts, cliente_folder, self.log)

            self.progress_bar.stop()
            self.progress_bar.config(mode="determinate")
            self.set_progress(100, "✅ Varredura concluída com sucesso!")
            self.log(f"[SUCESSO] Relatório gerado em: {filepath}")
            messagebox.showinfo("Sucesso", f"Varredura de hardware concluída!\n\nRelatório salvo em:\n{filepath}")
            self.check_files_in_current_folder()
        except Exception as e:
            self.progress_bar.stop()
            self.progress_bar.config(mode="determinate")
            self.set_progress(0, "❌ Falha na varredura.")
            self.log(f"[ERRO] Falha na varredura: {str(e)}")
            messagebox.showerror("Erro", str(e))
        finally:
            self.btn_coletar.config(state='normal')
            self.btn_run.config(state='normal')

    # --- EXECUÇÃO DE EXCEL/PDF EM THREAD ---

    def start_process_thread(self):
        threading.Thread(target=self.run_process, daemon=True).start()

    def run_process(self):
        cliente_folder = self.entry_cliente_folder.get().strip()
        excel_path = self.entry_excel_path.get().strip()
        pdf_path = os.path.splitext(excel_path)[0] + ".pdf"

        if not os.path.exists(cliente_folder):
            messagebox.showerror("Erro", "Pasta do cliente inválida.")
            return

        if not self.var_gen_excel.get() and not self.var_gen_pdf.get():
            messagebox.showwarning("Aviso", "Selecione pelo menos um formato de saída (Excel ou PDF).")
            return

        self.btn_coletar.config(state='disabled')
        self.btn_run.config(state='disabled')
        self.set_progress(0, "🚀 Iniciando consolidação de relatórios...")

        try:
            def cb_progress(pct):
                self.set_progress(pct, f"🚀 Processando dados... ({int(pct)}%)")

            if self.var_gen_excel.get():
                process_inventory(cliente_folder, excel_path, self.log, progress_callback=cb_progress)

            if self.var_gen_pdf.get():
                pdf_success = process_pdf_report(cliente_folder, pdf_path, self.log, progress_callback=cb_progress)
                if not pdf_success:
                    self.log("[AVISO] O relatório em PDF não pôde ser gerado. Verifique os logs acima.")

            self.set_progress(100, "✅ Processo concluído!")
            messagebox.showinfo("Sucesso", f"Processamento concluído!\n\nVerifique a pasta do cliente.")
        except Exception as e:
            self.set_progress(0, "❌ Erro ao exportar relatórios.")
            self.log(f"[ERRO] {str(e)}")
            messagebox.showerror("Erro Crítico", str(e))
        finally:
            self.btn_coletar.config(state='normal')
            self.btn_run.config(state='normal')

    def select_excel_destination(self):
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
            initialfile=os.path.basename(self.entry_excel_path.get())
        )
        if file_path:
            self.entry_excel_path.delete(0, tk.END)
            self.entry_excel_path.insert(0, file_path)

    def check_files_in_current_folder(self):
        folder = self.entry_cliente_folder.get().strip()
        if os.path.exists(folder):
            txts = glob.glob(os.path.join(folder, "*-info.txt"))
            self.log(f"[STATUS] Encontrado(s) {len(txts)} relatório(s) na pasta atual.")


if __name__ == "__main__":
    try:
        root = ttk.Window(themename="flatly")
    except Exception:
        root = tk.Tk()

    app = ExtractGUIApp(root)
    root.mainloop()